#!/usr/bin/env python3
"""
Generate a modern, aesthetically clean PowerPoint presentation with two slides:
1. 2026 Financials by ACT Pod
2. Frosty the Snow Man sheet music info
"""

from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.chart import XL_CHART_TYPE

# ── Colour palette ──────────────────────────────────────────────
WHITE       = RGBColor(0xFF, 0xFF, 0xFF)
OFF_WHITE   = RGBColor(0xF7, 0xF7, 0xFA)
LIGHT_GRAY  = RGBColor(0xE8, 0xE8, 0xED)
MID_GRAY    = RGBColor(0x9B, 0x9B, 0xA3)
DARK_GRAY   = RGBColor(0x4A, 0x4A, 0x55)
CHARCOAL    = RGBColor(0x2D, 0x2D, 0x3A)
NEAR_BLACK  = RGBColor(0x1A, 0x1A, 0x2E)
ACCENT_BLUE = RGBColor(0x3B, 0x82, 0xF6)
ACCENT_INDIGO = RGBColor(0x63, 0x66, 0xF1)
ACCENT_TEAL = RGBColor(0x14, 0xB8, 0xA6)
ACCENT_AMBER = RGBColor(0xF5, 0x9E, 0x0B)
ACCENT_ROSE = RGBColor(0xF4, 0x3F, 0x5E)
TABLE_HEADER_BG = RGBColor(0x1E, 0x1E, 0x2E)
TABLE_ROW_ALT   = RGBColor(0xF1, 0xF5, 0xF9)
TABLE_TOTAL_BG  = RGBColor(0xE0, 0xE7, 0xFF)
GRADIENT_START  = RGBColor(0x0F, 0x17, 0x2A)
GRADIENT_END    = RGBColor(0x1E, 0x29, 0x3B)

prs = Presentation()
prs.slide_width  = Inches(13.333)
prs.slide_height = Inches(7.5)

SLIDE_W = prs.slide_width
SLIDE_H = prs.slide_height

# ─── helper functions ───────────────────────────────────────────
def add_solid_bg(slide, color):
    bg = slide.background
    fill = bg.fill
    fill.solid()
    fill.fore_color.rgb = color

def add_shape(slide, left, top, width, height, fill_color, border_color=None, border_width=None, shape_type=MSO_SHAPE.ROUNDED_RECTANGLE):
    shape = slide.shapes.add_shape(shape_type, left, top, width, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill_color
    if border_color:
        shape.line.color.rgb = border_color
        shape.line.width = Pt(border_width or 1)
    else:
        shape.line.fill.background()
    shape.shadow.inherit = False
    return shape

def set_cell_text(cell, text, font_size=10, bold=False, color=CHARCOAL, alignment=PP_ALIGN.LEFT, font_name="Segoe UI"):
    cell.text = ""
    p = cell.text_frame.paragraphs[0]
    p.alignment = alignment
    run = p.add_run()
    run.text = str(text)
    run.font.size = Pt(font_size)
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.name = font_name
    cell.vertical_anchor = MSO_ANCHOR.MIDDLE
    cell.margin_left = Pt(8)
    cell.margin_right = Pt(8)
    cell.margin_top = Pt(4)
    cell.margin_bottom = Pt(4)

def set_cell_bg(cell, color):
    cell.fill.solid()
    cell.fill.fore_color.rgb = color

def add_text_box(slide, left, top, width, height, text, font_size=12,
                 bold=False, color=CHARCOAL, alignment=PP_ALIGN.LEFT,
                 font_name="Segoe UI"):
    txBox = slide.shapes.add_textbox(left, top, width, height)
    tf = txBox.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = alignment
    run = p.add_run()
    run.text = text
    run.font.size = Pt(font_size)
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.name = font_name
    return txBox

def add_accent_line(slide, left, top, width, color=ACCENT_BLUE, height=Pt(3)):
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, width, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.fill.background()
    shape.shadow.inherit = False
    return shape


# ═══════════════════════════════════════════════════════════════
# SLIDE 1 – 2026 Financials by ACT Pod
# ═══════════════════════════════════════════════════════════════
slide1 = prs.slides.add_slide(prs.slide_layouts[6])  # blank
add_solid_bg(slide1, WHITE)

# ── Top accent bar ──
add_shape(slide1, Inches(0), Inches(0), SLIDE_W, Pt(5), ACCENT_BLUE, shape_type=MSO_SHAPE.RECTANGLE)

# ── Title section ──
add_text_box(slide1, Inches(0.8), Inches(0.45), Inches(8), Inches(0.55),
             "2026 Financials by ACT Pod", font_size=28, bold=True,
             color=NEAR_BLACK, font_name="Segoe UI Semibold")

add_accent_line(slide1, Inches(0.8), Inches(1.05), Inches(2.2), ACCENT_BLUE, Pt(3))

add_text_box(slide1, Inches(0.8), Inches(1.2), Inches(6), Inches(0.35),
             "Revenue & Growth Planning  |  US/GCSO Division", font_size=11,
             color=MID_GRAY, font_name="Segoe UI")

# ── Financial Table ──
rows, cols = 13, 4
tbl_left = Inches(0.8)
tbl_top  = Inches(1.8)
tbl_w    = Inches(7.0)
tbl_h    = Inches(5.0)

table_shape = slide1.shapes.add_table(rows, cols, tbl_left, tbl_top, tbl_w, tbl_h)
table = table_shape.table

# Column widths
col_widths = [Inches(2.6), Inches(1.5), Inches(1.5), Inches(1.4)]
for i, w in enumerate(col_widths):
    table.columns[i].width = w

# Header row
headers = ["$ MM", "2025 Revenue", "2026 Plan", "Planned Growth"]
for j, h in enumerate(headers):
    set_cell_text(table.cell(0, j), h, font_size=10, bold=True, color=WHITE,
                  alignment=PP_ALIGN.CENTER, font_name="Segoe UI Semibold")
    set_cell_bg(table.cell(0, j), TABLE_HEADER_BG)

# Data
data = [
    ("US/GCSO_HCP",              "$23.5", "",       ""),
    ("US/GCSO_B2B",              "$10.8", "",       ""),
    ("US/GCSO_Patient",          "$7.1",  "",       ""),
    ("US/GCSO_Medical",          "$3.4",  "",       ""),
    ("US/GCSO_Field_GTM_SFE",    "$14.0", "",       ""),
    ("US/GCSO_Financial_Planning","$7.5", "",       ""),
    ("US/GCSO_Brand/BU_Advisory","$8.0",  "",       ""),
    ("US/GCSO_Customer_Insights","$3.9",  "",       ""),
    ("US/GCSO_Scaled_Analytics", "$14.5", "",       ""),
    ("US/GCSO_Data_Digital/AI",  "$5.4",  "",       ""),
    ("US/GSCO Total",            "$98.0", "$101.0", ""),
]

for i, (name, rev, plan, growth) in enumerate(data):
    row_idx = i + 1
    is_total = (i == len(data) - 1)
    is_alt   = (i % 2 == 1 and not is_total)

    bg = TABLE_TOTAL_BG if is_total else (TABLE_ROW_ALT if is_alt else WHITE)

    for j, val in enumerate([name, rev, plan, growth]):
        cell = table.cell(row_idx, j)
        set_cell_bg(cell, bg)
        align = PP_ALIGN.LEFT if j == 0 else PP_ALIGN.CENTER
        set_cell_text(cell, val, font_size=10,
                      bold=is_total,
                      color=NEAR_BLACK if is_total else DARK_GRAY,
                      alignment=align)

    # Style row height
    table.rows[row_idx].height = Inches(0.36)

table.rows[0].height = Inches(0.42)

# Remove table borders and add subtle lines
for row in table.rows:
    for cell in row.cells:
        for border_tag in ['a:lnL', 'a:lnR', 'a:lnT', 'a:lnB']:
            from lxml import etree
            tcPr = cell._tc.get_or_add_tcPr()
            nsmap = {'a': 'http://schemas.openxmlformats.org/drawingml/2006/main'}
            for ln in tcPr.findall(border_tag, nsmap):
                tcPr.remove(ln)
            ln_elem = etree.SubElement(tcPr, f'{{{nsmap["a"]}}}{border_tag[2:]}')
            ln_elem.set('w', '6350')  # 0.5pt in EMUs
            solid = etree.SubElement(ln_elem, f'{{{nsmap["a"]}}}solidFill')
            srgb = etree.SubElement(solid, f'{{{nsmap["a"]}}}srgbClr')
            srgb.set('val', 'E2E8F0')

# ── Right panel – Info card ──
card_left  = Inches(8.3)
card_top   = Inches(1.8)
card_w     = Inches(4.4)
card_h     = Inches(5.0)

card = add_shape(slide1, card_left, card_top, card_w, card_h, OFF_WHITE,
                 border_color=LIGHT_GRAY, border_width=1)

# Card accent strip
add_shape(slide1, card_left, card_top, Pt(4), card_h, ACCENT_INDIGO, shape_type=MSO_SHAPE.RECTANGLE)

# Card title
add_text_box(slide1, Inches(8.6), Inches(2.0), Inches(3.9), Inches(0.4),
             "Pod-Level Reporting", font_size=14, bold=True,
             color=NEAR_BLACK, font_name="Segoe UI Semibold")

add_accent_line(slide1, Inches(8.6), Inches(2.4), Inches(1.4), ACCENT_INDIGO, Pt(2))

# Card subtitle
add_text_box(slide1, Inches(8.6), Inches(2.55), Inches(3.9), Inches(0.65),
             "ACT Pod leaders will receive Pod level P&L and BD report "
             "throughout the year sent via email by FBPs covering:",
             font_size=9.5, color=DARK_GRAY)

# Bullet items
bullets = [
    "YTD Actuals compared to Plan",
    "At-Risk Projects",
    "YTD Committed Spend",
    "BD Analysis + Cost Breakdown",
    "Work Type Portfolio mix by delivery margin",
]

bullet_y = Inches(3.3)
for b in bullets:
    txBox = slide1.shapes.add_textbox(Inches(8.6), bullet_y, Inches(3.9), Inches(0.3))
    tf = txBox.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.LEFT
    run = p.add_run()
    run.text = f"\u2022  {b}"
    run.font.size = Pt(9.5)
    run.font.color.rgb = DARK_GRAY
    run.font.name = "Segoe UI"
    bullet_y += Inches(0.28)

# Expectations section
add_text_box(slide1, Inches(8.6), Inches(4.85), Inches(3.9), Inches(0.35),
             "ACT POD Leader Expectations", font_size=12, bold=True,
             color=NEAR_BLACK, font_name="Segoe UI Semibold")

add_accent_line(slide1, Inches(8.6), Inches(5.2), Inches(1.2), ACCENT_TEAL, Pt(2))

add_text_box(slide1, Inches(8.6), Inches(5.35), Inches(3.9), Inches(0.6),
             "Produce a one-page quarterly objectives report with defined "
             "targets and ongoing status updates.",
             font_size=9.5, color=DARK_GRAY)

# Footer note
add_text_box(slide1, Inches(0.8), Inches(6.9), Inches(7), Inches(0.35),
             "Please share 2026 revenue and growth plan with Emma",
             font_size=9, color=MID_GRAY, font_name="Segoe UI")

# Footer right - page number
add_text_box(slide1, Inches(11.5), Inches(7.0), Inches(1.5), Inches(0.3),
             "7", font_size=9, color=MID_GRAY, alignment=PP_ALIGN.RIGHT)

# Confidential badge
conf_shape = add_shape(slide1, Inches(11.0), Inches(6.85), Inches(1.6), Inches(0.3),
                        OFF_WHITE, border_color=LIGHT_GRAY, border_width=0.5)
add_text_box(slide1, Inches(11.0), Inches(6.87), Inches(1.6), Inches(0.28),
             "Confidential", font_size=8, color=MID_GRAY,
             alignment=PP_ALIGN.CENTER, font_name="Segoe UI")


# ═══════════════════════════════════════════════════════════════
# SLIDE 2 – Frosty the Snow Man
# ═══════════════════════════════════════════════════════════════
slide2 = prs.slides.add_slide(prs.slide_layouts[6])
add_solid_bg(slide2, WHITE)

# Top accent bar
add_shape(slide2, Inches(0), Inches(0), SLIDE_W, Pt(5), ACCENT_ROSE, shape_type=MSO_SHAPE.RECTANGLE)

# Title
add_text_box(slide2, Inches(0.8), Inches(0.45), Inches(8), Inches(0.55),
             "Frosty the Snow Man", font_size=28, bold=True,
             color=NEAR_BLACK, font_name="Segoe UI Semibold")

add_accent_line(slide2, Inches(0.8), Inches(1.05), Inches(2.2), ACCENT_ROSE, Pt(3))

add_text_box(slide2, Inches(0.8), Inches(1.2), Inches(8), Inches(0.35),
             "Words and Music by Steve Nelson and Jack Rollins  |  Arr. by Tom Gerou",
             font_size=11, color=MID_GRAY, font_name="Segoe UI")

# ── Two-column layout ──

# LEFT COLUMN — Song details card
left_card = add_shape(slide2, Inches(0.8), Inches(1.85), Inches(5.8), Inches(5.0),
                       OFF_WHITE, border_color=LIGHT_GRAY, border_width=1)
add_shape(slide2, Inches(0.8), Inches(1.85), Pt(4), Inches(5.0), ACCENT_ROSE, shape_type=MSO_SHAPE.RECTANGLE)

add_text_box(slide2, Inches(1.1), Inches(2.0), Inches(5.2), Inches(0.35),
             "Song Overview", font_size=14, bold=True,
             color=NEAR_BLACK, font_name="Segoe UI Semibold")
add_accent_line(slide2, Inches(1.1), Inches(2.38), Inches(1.2), ACCENT_ROSE, Pt(2))

# Details table
detail_rows = 7
detail_cols = 2
dtbl_shape = slide2.shapes.add_table(detail_rows, detail_cols, Inches(1.1), Inches(2.6),
                                       Inches(5.2), Inches(2.8))
dtbl = dtbl_shape.table
dtbl.columns[0].width = Inches(1.8)
dtbl.columns[1].width = Inches(3.4)

details = [
    ("Tempo",       "Allegro"),
    ("Key",         "C Major"),
    ("Time Signature", "4/4 (Common Time)"),
    ("Measures",    "40 measures"),
    ("Difficulty",  "Early Intermediate Piano"),
    ("Copyright",   "\u00a9 1950 by Chappell & Co., Copyright Renewed"),
    ("Features",    "Both hands reading, finger numbering, dynamic markings (mp)"),
]

for i, (label, value) in enumerate(details):
    bg = TABLE_ROW_ALT if i % 2 == 0 else WHITE
    set_cell_bg(dtbl.cell(i, 0), bg)
    set_cell_bg(dtbl.cell(i, 1), bg)
    set_cell_text(dtbl.cell(i, 0), label, font_size=10, bold=True,
                  color=NEAR_BLACK, font_name="Segoe UI Semibold")
    set_cell_text(dtbl.cell(i, 1), value, font_size=10, color=DARK_GRAY)
    dtbl.rows[i].height = Inches(0.38)

# Remove harsh borders from detail table
for row in dtbl.rows:
    for cell in row.cells:
        for border_tag in ['a:lnL', 'a:lnR', 'a:lnT', 'a:lnB']:
            tcPr = cell._tc.get_or_add_tcPr()
            nsmap = {'a': 'http://schemas.openxmlformats.org/drawingml/2006/main'}
            for ln in tcPr.findall(border_tag, nsmap):
                tcPr.remove(ln)
            ln_elem = etree.SubElement(tcPr, f'{{{nsmap["a"]}}}{border_tag[2:]}')
            ln_elem.set('w', '6350')
            solid = etree.SubElement(ln_elem, f'{{{nsmap["a"]}}}solidFill')
            srgb = etree.SubElement(solid, f'{{{nsmap["a"]}}}srgbClr')
            srgb.set('val', 'E2E8F0')

# Notable directions
add_text_box(slide2, Inches(1.1), Inches(5.6), Inches(5.2), Inches(0.3),
             "Performance Notes", font_size=11, bold=True,
             color=NEAR_BLACK, font_name="Segoe UI Semibold")
add_text_box(slide2, Inches(1.1), Inches(5.9), Inches(5.2), Inches(0.8),
             "\u2022  Measures 33-end: Both hands 8va to end\n"
             "\u2022  Dynamic marking: mp (mezzo piano) at measure 33\n"
             "\u2022  Ending passage: \"Thump-et-y thump thump, over the hills of snow\"",
             font_size=9.5, color=DARK_GRAY)

# RIGHT COLUMN — Lyrics card
right_card = add_shape(slide2, Inches(7.0), Inches(1.85), Inches(5.6), Inches(5.0),
                        OFF_WHITE, border_color=LIGHT_GRAY, border_width=1)
add_shape(slide2, Inches(7.0), Inches(1.85), Pt(4), Inches(5.0), ACCENT_AMBER, shape_type=MSO_SHAPE.RECTANGLE)

add_text_box(slide2, Inches(7.3), Inches(2.0), Inches(5.0), Inches(0.35),
             "Lyrics", font_size=14, bold=True,
             color=NEAR_BLACK, font_name="Segoe UI Semibold")
add_accent_line(slide2, Inches(7.3), Inches(2.38), Inches(0.8), ACCENT_AMBER, Pt(2))

# Verse 1
add_text_box(slide2, Inches(7.3), Inches(2.55), Inches(2.2), Inches(0.25),
             "Verse 1", font_size=10, bold=True, color=ACCENT_ROSE, font_name="Segoe UI Semibold")

verse1 = (
    "Frosty the Snow Man was a jolly happy soul,\n"
    "With a corn cob pipe and a button nose\n"
    "and two eyes made out of coal.\n\n"
    "Frosty the Snow Man is a fairy tale, they say,\n"
    "He was made of snow but the children know\n"
    "how he came to life one day.\n\n"
    "There must have been some magic\n"
    "in that old silk hat they found.\n"
    "For when they placed it on his head\n"
    "he began to dance around."
)
add_text_box(slide2, Inches(7.3), Inches(2.8), Inches(5.0), Inches(2.2),
             verse1, font_size=9, color=DARK_GRAY, font_name="Segoe UI")

# Verse 2
add_text_box(slide2, Inches(7.3), Inches(5.0), Inches(2.2), Inches(0.25),
             "Verse 2", font_size=10, bold=True, color=ACCENT_ROSE, font_name="Segoe UI Semibold")

verse2 = (
    "Frosty the Snow Man knew the sun was hot that day,\n"
    "So he said, \"Let's run and we'll have some fun\n"
    "now before I melt away.\"\n\n"
    "Frosty the Snow Man had to hurry on his way,\n"
    "But he waved goodbye sayin' \"Don't you cry,\n"
    "I'll be back again some day.\""
)
add_text_box(slide2, Inches(7.3), Inches(5.25), Inches(5.0), Inches(1.5),
             verse2, font_size=9, color=DARK_GRAY, font_name="Segoe UI")

# Footer
add_text_box(slide2, Inches(11.0), Inches(6.87), Inches(1.6), Inches(0.28),
             "Page 16-17", font_size=8, color=MID_GRAY,
             alignment=PP_ALIGN.CENTER, font_name="Segoe UI")


# ═══════════════════════════════════════════════════════════════
# Save
# ═══════════════════════════════════════════════════════════════
output_path = "/home/user/airmsg/Modern_Slides.pptx"
prs.save(output_path)
print(f"Presentation saved to {output_path}")
