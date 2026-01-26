#!/usr/bin/env python3
"""
Create modern, clean executive PowerPoint slides for MarComms Operating Model Redesign
"""

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

# Modern color palette
COLORS = {
    'primary': RGBColor(0, 82, 147),       # Deep blue
    'secondary': RGBColor(0, 150, 199),    # Teal blue
    'accent': RGBColor(255, 107, 53),      # Coral orange
    'dark': RGBColor(51, 51, 51),          # Dark gray for text
    'medium': RGBColor(102, 102, 102),     # Medium gray
    'light': RGBColor(240, 240, 240),      # Light gray background
    'white': RGBColor(255, 255, 255),      # White
    'success': RGBColor(46, 184, 92),      # Green
}

def set_font(run, size=12, bold=False, color=None, font_name='Segoe UI'):
    """Set font properties for a text run"""
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.name = font_name
    if color:
        run.font.color.rgb = color

def add_title_slide(prs, title, subtitle=None):
    """Add a title slide"""
    slide_layout = prs.slide_layouts[6]  # Blank layout
    slide = prs.slides.add_slide(slide_layout)

    # Add accent bar at top
    bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.33), Inches(0.15))
    bar.fill.solid()
    bar.fill.fore_color.rgb = COLORS['primary']
    bar.line.fill.background()

    # Title
    title_box = slide.shapes.add_textbox(Inches(0.5), Inches(2.5), Inches(12.33), Inches(1))
    tf = title_box.text_frame
    p = tf.paragraphs[0]
    p.text = title
    p.alignment = PP_ALIGN.CENTER
    set_font(p.runs[0], size=44, bold=True, color=COLORS['primary'])

    # Subtitle
    if subtitle:
        sub_box = slide.shapes.add_textbox(Inches(0.5), Inches(3.6), Inches(12.33), Inches(0.6))
        tf = sub_box.text_frame
        p = tf.paragraphs[0]
        p.text = subtitle
        p.alignment = PP_ALIGN.CENTER
        set_font(p.runs[0], size=20, color=COLORS['medium'])

    return slide

def add_content_slide(prs, title):
    """Add a content slide with title"""
    slide_layout = prs.slide_layouts[6]  # Blank layout
    slide = prs.slides.add_slide(slide_layout)

    # Add accent bar at top
    bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.33), Inches(0.08))
    bar.fill.solid()
    bar.fill.fore_color.rgb = COLORS['primary']
    bar.line.fill.background()

    # Title
    title_box = slide.shapes.add_textbox(Inches(0.5), Inches(0.3), Inches(12.33), Inches(0.7))
    tf = title_box.text_frame
    p = tf.paragraphs[0]
    p.text = title
    set_font(p.runs[0], size=28, bold=True, color=COLORS['primary'])

    return slide

def add_bullet_box(slide, left, top, width, height, title, bullets, title_color=None, bg_color=None):
    """Add a box with title and bullet points"""
    if title_color is None:
        title_color = COLORS['primary']

    # Background shape if needed
    if bg_color:
        bg = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
        bg.fill.solid()
        bg.fill.fore_color.rgb = bg_color
        bg.line.fill.background()

    # Title
    title_box = slide.shapes.add_textbox(left + Inches(0.15), top + Inches(0.1), width - Inches(0.3), Inches(0.4))
    tf = title_box.text_frame
    p = tf.paragraphs[0]
    p.text = title
    set_font(p.runs[0], size=14, bold=True, color=title_color)

    # Bullets
    bullet_box = slide.shapes.add_textbox(left + Inches(0.15), top + Inches(0.45), width - Inches(0.3), height - Inches(0.55))
    tf = bullet_box.text_frame
    tf.word_wrap = True

    for i, bullet in enumerate(bullets):
        if i == 0:
            p = tf.paragraphs[0]
        else:
            p = tf.add_paragraph()
        p.text = f"• {bullet}"
        p.space_before = Pt(4)
        p.space_after = Pt(2)
        set_font(p.runs[0], size=11, color=COLORS['dark'])

def add_key_insight_box(slide, left, top, width, text, subtext=None):
    """Add a key insight/bottom line box"""
    # Background
    bg = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, Inches(0.9))
    bg.fill.solid()
    bg.fill.fore_color.rgb = COLORS['light']
    bg.line.color.rgb = COLORS['primary']
    bg.line.width = Pt(2)

    # Label
    label_box = slide.shapes.add_textbox(left + Inches(0.2), top + Inches(0.1), Inches(1.5), Inches(0.3))
    tf = label_box.text_frame
    p = tf.paragraphs[0]
    p.text = "KEY INSIGHT"
    set_font(p.runs[0], size=10, bold=True, color=COLORS['accent'])

    # Main text
    text_box = slide.shapes.add_textbox(left + Inches(0.2), top + Inches(0.35), width - Inches(0.4), Inches(0.5))
    tf = text_box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = text
    set_font(p.runs[0], size=12, bold=True, color=COLORS['dark'])

    if subtext:
        p = tf.add_paragraph()
        p.text = subtext
        set_font(p.runs[0], size=11, color=COLORS['medium'])

def create_slide_1(prs):
    """Slide 1: What We Have Learned"""
    slide = add_content_slide(prs, "Background: What We Have Learned")

    # Left column - Discovery Conversations
    add_bullet_box(
        slide,
        Inches(0.4), Inches(1.1), Inches(6.1), Inches(2.7),
        "From Discovery Conversations",
        [
            "MarComms operates in a complex, high-volume, matrixed environment",
            "Most work requires multiple teams and handoffs",
            "Review cycles are heavy and timelines remain long",
            "Ownership and visibility are not always clear",
            "Speed and integration are recurring challenges",
            "Teams are highly committed to quality and reputation"
        ],
        title_color=COLORS['secondary'],
        bg_color=COLORS['light']
    )

    # Right column - Past Efforts
    add_bullet_box(
        slide,
        Inches(6.8), Inches(1.1), Inches(6.1), Inches(2.7),
        "From Past Improvement Efforts",
        [
            "This problem has been addressed multiple times before",
            "Some timelines improved initially",
            "Improvements were not sustained over time",
            "Many outputs remain slower than external benchmarks",
            "Deeper, systemic change has been hard to achieve"
        ],
        title_color=COLORS['secondary'],
        bg_color=COLORS['light']
    )

    # Bottom line box
    add_key_insight_box(
        slide,
        Inches(0.4), Inches(4.0), Inches(12.5),
        "We have made progress before, but it has not scaled or lasted."
    )

    # Footer
    footer = slide.shapes.add_textbox(Inches(0.4), Inches(5.1), Inches(12), Inches(0.3))
    tf = footer.text_frame
    p = tf.paragraphs[0]
    p.text = "Based on discovery sessions with key stakeholders and functional leaders"
    set_font(p.runs[0], size=9, color=COLORS['medium'])

    return slide

def create_slide_2(prs):
    """Slide 2: Why This Is Hard to Fix"""
    slide = add_content_slide(prs, "Why This Has Been Hard to Fix")

    # Four columns
    col_width = Inches(3.0)
    col_height = Inches(2.8)
    start_x = Inches(0.4)
    gap = Inches(0.15)
    top_y = Inches(1.0)

    columns = [
        ("Process & Structure", COLORS['primary'], [
            "Existing workflows are difficult to disrupt",
            "End-to-end ownership is unclear",
            "Different content types treated similarly",
            "Processes reflect history more than current needs"
        ]),
        ("Culture & Behavior", COLORS['secondary'], [
            "Strong emphasis on perfection over speed",
            "Risk aversion driven by past experiences",
            "\"More work\" sometimes seen as \"more value\"",
            "Many handoffs treated as individual jobs"
        ]),
        ("Change & Adoption", COLORS['accent'], [
            "Past efforts focused on process, not behavior",
            "Change management was limited",
            "Stakeholder experience not always central",
            "Alignment on urgency has varied"
        ]),
        ("External Pressure", COLORS['success'], [
            "Faster-moving competitors",
            "New digital and AI channels emerging",
            "Changing content consumption patterns",
            "Rising expectations for responsiveness"
        ])
    ]

    for i, (title, color, bullets) in enumerate(columns):
        x = start_x + (col_width + gap) * i

        # Header bar
        header = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, top_y, col_width, Inches(0.35))
        header.fill.solid()
        header.fill.fore_color.rgb = color
        header.line.fill.background()

        # Header text
        header_text = slide.shapes.add_textbox(x, top_y + Inches(0.05), col_width, Inches(0.3))
        tf = header_text.text_frame
        p = tf.paragraphs[0]
        p.text = title
        p.alignment = PP_ALIGN.CENTER
        set_font(p.runs[0], size=12, bold=True, color=COLORS['white'])

        # Bullet content
        bullet_box = slide.shapes.add_textbox(x + Inches(0.1), top_y + Inches(0.45), col_width - Inches(0.2), col_height - Inches(0.5))
        tf = bullet_box.text_frame
        tf.word_wrap = True

        for j, bullet in enumerate(bullets):
            if j == 0:
                p = tf.paragraphs[0]
            else:
                p = tf.add_paragraph()
            p.text = f"• {bullet}"
            p.space_before = Pt(6)
            p.space_after = Pt(2)
            set_font(p.runs[0], size=10, color=COLORS['dark'])

    # Key insight box at bottom
    insight_top = Inches(4.0)
    bg = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.4), insight_top, Inches(12.5), Inches(0.85))
    bg.fill.solid()
    bg.fill.fore_color.rgb = COLORS['light']
    bg.line.color.rgb = COLORS['primary']
    bg.line.width = Pt(2)

    label_box = slide.shapes.add_textbox(Inches(0.6), insight_top + Inches(0.08), Inches(1.5), Inches(0.25))
    tf = label_box.text_frame
    p = tf.paragraphs[0]
    p.text = "KEY INSIGHT"
    set_font(p.runs[0], size=10, bold=True, color=COLORS['accent'])

    text_box = slide.shapes.add_textbox(Inches(0.6), insight_top + Inches(0.32), Inches(11.9), Inches(0.5))
    tf = text_box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "This is not just a workflow problem. It is a system-level challenge spanning structure, culture, and adoption."
    set_font(p.runs[0], size=12, bold=True, color=COLORS['dark'])
    p = tf.add_paragraph()
    p.text = "Sustainable improvement requires coordinated redesign across all four dimensions."
    set_font(p.runs[0], size=11, color=COLORS['medium'])

    return slide

def create_slide_3(prs):
    """Slide 3: Project Plan Overview"""
    slide = add_content_slide(prs, "Project Plan: MarComms Operating Model Redesign")

    # Subtitle
    sub = slide.shapes.add_textbox(Inches(0.5), Inches(0.65), Inches(12), Inches(0.3))
    tf = sub.text_frame
    p = tf.paragraphs[0]
    p.text = "This work starts with understanding why prior attempts stalled — not with process mapping"
    set_font(p.runs[0], size=12, color=COLORS['medium'])

    # Four phase boxes
    phases = [
        ("Phase 1", "Deepen Problem Understanding", "Weeks 1–4", COLORS['primary'], [
            "Discovery sessions with key stakeholders",
            "Review past improvement efforts",
            "Identify true bottlenecks vs. symptoms",
            "Segment work types by risk and complexity"
        ]),
        ("Phase 2", "Prioritize What to Reimagine", "Weeks 5–6", COLORS['secondary'], [
            "Identify priority areas by impact",
            "Decide what to redesign now vs. later",
            "Align on design principles",
            "Define what \"good enough\" means"
        ]),
        ("Phase 3", "Reimagine Operating Model", "Weeks 7–11", COLORS['accent'], [
            "Redesign priority workflows end-to-end",
            "Clarify ownership and decision rights",
            "Build differentiation into the system",
            "Simplify handoffs and approvals"
        ]),
        ("Phase 4", "Change & Adoption", "Weeks 5–12", COLORS['success'], [
            "Stakeholder engagement throughout",
            "Leadership alignment on behaviors",
            "Communication and enablement plan",
            "Define success measures"
        ])
    ]

    box_width = Inches(3.0)
    box_height = Inches(2.6)
    start_x = Inches(0.4)
    gap = Inches(0.15)
    top_y = Inches(1.05)

    for i, (phase, title, timing, color, bullets) in enumerate(phases):
        x = start_x + (box_width + gap) * i

        # Phase header
        header = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, top_y, box_width, Inches(0.6))
        header.fill.solid()
        header.fill.fore_color.rgb = color
        header.line.fill.background()

        # Phase number and title
        phase_box = slide.shapes.add_textbox(x, top_y + Inches(0.05), box_width, Inches(0.25))
        tf = phase_box.text_frame
        p = tf.paragraphs[0]
        p.text = phase
        p.alignment = PP_ALIGN.CENTER
        set_font(p.runs[0], size=10, bold=True, color=COLORS['white'])

        title_box = slide.shapes.add_textbox(x, top_y + Inches(0.28), box_width, Inches(0.25))
        tf = title_box.text_frame
        p = tf.paragraphs[0]
        p.text = title
        p.alignment = PP_ALIGN.CENTER
        set_font(p.runs[0], size=11, bold=True, color=COLORS['white'])

        # Timing badge
        timing_box = slide.shapes.add_textbox(x, top_y + Inches(0.65), box_width, Inches(0.25))
        tf = timing_box.text_frame
        p = tf.paragraphs[0]
        p.text = timing
        p.alignment = PP_ALIGN.CENTER
        set_font(p.runs[0], size=10, bold=True, color=color)

        # Bullets
        bullet_box = slide.shapes.add_textbox(x + Inches(0.1), top_y + Inches(0.95), box_width - Inches(0.2), box_height - Inches(1.0))
        tf = bullet_box.text_frame
        tf.word_wrap = True

        for j, bullet in enumerate(bullets):
            if j == 0:
                p = tf.paragraphs[0]
            else:
                p = tf.add_paragraph()
            p.text = f"• {bullet}"
            p.space_before = Pt(5)
            p.space_after = Pt(2)
            set_font(p.runs[0], size=10, color=COLORS['dark'])

    # Guiding principles at bottom
    principles_top = Inches(3.85)

    # Label
    label = slide.shapes.add_textbox(Inches(0.4), principles_top, Inches(2), Inches(0.3))
    tf = label.text_frame
    p = tf.paragraphs[0]
    p.text = "GUIDING PRINCIPLES"
    set_font(p.runs[0], size=10, bold=True, color=COLORS['primary'])

    # Principle boxes
    principles = [
        "Redesign before optimize",
        "Focus on impact, not volume",
        "Differentiate by risk",
        "Build for adoption"
    ]

    pill_start = Inches(0.4)
    pill_width = Inches(2.9)
    pill_gap = Inches(0.2)

    for i, principle in enumerate(principles):
        x = pill_start + (pill_width + pill_gap) * i

        pill = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, principles_top + Inches(0.35), pill_width, Inches(0.4))
        pill.fill.solid()
        pill.fill.fore_color.rgb = COLORS['light']
        pill.line.color.rgb = COLORS['primary']
        pill.line.width = Pt(1)

        text_box = slide.shapes.add_textbox(x, principles_top + Inches(0.42), pill_width, Inches(0.3))
        tf = text_box.text_frame
        p = tf.paragraphs[0]
        p.text = principle
        p.alignment = PP_ALIGN.CENTER
        set_font(p.runs[0], size=11, bold=True, color=COLORS['primary'])

    return slide

def create_slide_4(prs):
    """Slide 4: Timeline at a Glance"""
    slide = add_content_slide(prs, "Timeline at a Glance")

    # Timeline visual
    timeline_top = Inches(1.2)
    timeline_height = Inches(2.8)

    # Week markers
    weeks_box = slide.shapes.add_textbox(Inches(1.8), timeline_top, Inches(10.5), Inches(0.3))
    tf = weeks_box.text_frame
    p = tf.paragraphs[0]
    p.text = "Week 1        2        3        4        5        6        7        8        9        10       11       12"
    set_font(p.runs[0], size=10, color=COLORS['medium'], font_name='Consolas')

    # Phase bars
    phases = [
        ("Phase 1: Deep Problem Discovery", 0, 4, COLORS['primary']),
        ("Phase 2: Prioritization & Focus", 4, 2, COLORS['secondary']),
        ("Phase 3: Operating Model Redesign", 6, 5, COLORS['accent']),
        ("Phase 4: Change & Adoption", 4, 8, COLORS['success']),
    ]

    bar_start_x = Inches(1.8)
    week_width = Inches(0.85)
    bar_height = Inches(0.5)
    bar_gap = Inches(0.15)

    for i, (label, start_week, duration, color) in enumerate(phases):
        y = timeline_top + Inches(0.5) + (bar_height + bar_gap) * i
        x = bar_start_x + (week_width * start_week)
        width = week_width * duration

        # Label on left
        label_box = slide.shapes.add_textbox(Inches(0.3), y + Inches(0.1), Inches(1.4), Inches(0.3))
        tf = label_box.text_frame
        p = tf.paragraphs[0]
        p.text = f"Phase {i+1}"
        set_font(p.runs[0], size=11, bold=True, color=color)

        # Bar
        bar = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, width, bar_height)
        bar.fill.solid()
        bar.fill.fore_color.rgb = color
        bar.line.fill.background()

        # Bar text
        bar_text = slide.shapes.add_textbox(x + Inches(0.1), y + Inches(0.12), width - Inches(0.2), Inches(0.3))
        tf = bar_text.text_frame
        p = tf.paragraphs[0]
        # Truncate long labels for bar
        short_labels = {
            "Phase 1: Deep Problem Discovery": "Deep Problem Discovery",
            "Phase 2: Prioritization & Focus": "Prioritization & Focus",
            "Phase 3: Operating Model Redesign": "Operating Model Redesign",
            "Phase 4: Change & Adoption": "Change Management & Adoption"
        }
        p.text = short_labels.get(label, label)
        set_font(p.runs[0], size=10, bold=True, color=COLORS['white'])

    # Summary table below timeline
    table_top = Inches(3.4)

    # Table header
    header_data = ["Phase", "Focus", "Timing", "Key Outputs"]
    col_widths = [Inches(1.3), Inches(4.0), Inches(1.5), Inches(5.5)]

    # Header row
    x_pos = Inches(0.4)
    for j, (header, width) in enumerate(zip(header_data, col_widths)):
        cell = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, x_pos, table_top, width, Inches(0.35))
        cell.fill.solid()
        cell.fill.fore_color.rgb = COLORS['primary']
        cell.line.color.rgb = COLORS['white']
        cell.line.width = Pt(1)

        text = slide.shapes.add_textbox(x_pos + Inches(0.1), table_top + Inches(0.05), width - Inches(0.2), Inches(0.3))
        tf = text.text_frame
        p = tf.paragraphs[0]
        p.text = header
        set_font(p.runs[0], size=10, bold=True, color=COLORS['white'])
        x_pos += width

    # Data rows
    table_data = [
        ("1", "Deep problem discovery", "Weeks 1–4", "Root cause analysis, evidence-based problem statement"),
        ("2", "Prioritization & focus", "Weeks 5–6", "Prioritized scope, executive alignment"),
        ("3", "Operating model redesign", "Weeks 7–11", "Future-state model, simplified workflows"),
        ("4", "Change & adoption", "Weeks 5–12", "Change strategy, adoption roadmap, success measures"),
    ]

    row_colors = [COLORS['primary'], COLORS['secondary'], COLORS['accent'], COLORS['success']]

    for i, (phase, focus, timing, outputs) in enumerate(table_data):
        y = table_top + Inches(0.35) + (Inches(0.4) * i)
        x_pos = Inches(0.4)
        row_data = [phase, focus, timing, outputs]

        for j, (data, width) in enumerate(zip(row_data, col_widths)):
            bg_color = COLORS['light'] if i % 2 == 0 else COLORS['white']
            cell = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, x_pos, y, width, Inches(0.4))
            cell.fill.solid()
            cell.fill.fore_color.rgb = bg_color
            cell.line.color.rgb = COLORS['light']
            cell.line.width = Pt(1)

            text = slide.shapes.add_textbox(x_pos + Inches(0.1), y + Inches(0.08), width - Inches(0.2), Inches(0.3))
            tf = text.text_frame
            tf.word_wrap = True
            p = tf.paragraphs[0]
            p.text = data
            text_color = row_colors[i] if j == 0 else COLORS['dark']
            set_font(p.runs[0], size=10, bold=(j==0), color=text_color)
            x_pos += width

    return slide

def main():
    # Create presentation with widescreen dimensions
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    # Create slides
    add_title_slide(prs, "MarComms Operating Model Redesign", "Executive Background & Project Plan")
    create_slide_1(prs)
    create_slide_2(prs)
    create_slide_3(prs)
    create_slide_4(prs)

    # Save
    output_path = "/home/user/airmsg/MarComms_Executive_Slides.pptx"
    prs.save(output_path)
    print(f"Presentation saved to: {output_path}")

if __name__ == "__main__":
    main()
